from pathlib import Path

from afs_assistant.config import Settings
from afs_assistant.core.models import SourceVersion, SourceMaterial
from afs_assistant.indexing.official import OfficialExtractor, OfficialChunker
from afs_assistant.indexing.references import extract_references, resolve_supplied_references, bibliography_element


def source(sid="S01", **extra):
    return SourceVersion(source_id=sid, sha256="a"*64, title="Test", uri="local", filename="test.pdf",
        kind="guidance", expected_pages=1, authority="provided", **extra)


def test_bibliography_resolves_only_verified_supplied_identity():
    records = extract_references(source(), ["Text\nReferences:\n- Author. A study. https://doi.org/10.1021/acs.analchem.1c00080\n\n2. Next section\nActual body."])
    assert len(records) == 1
    assert "Next section" not in records[0].citation
    target = source("S04", citation_keys=["10.1021/acs.analchem.1c00080"])
    resolve_supplied_references(records, [source(), target])
    assert records[0].target_source_id == "S04"
    assert bibliography_element(records[0].citation, records)
    assert not bibliography_element("Measurements differed from the earlier study [1].", records)


def test_short_bibliography_entries_are_not_treated_as_prose():
    records = extract_references(source("S04"), ["REFERENCES\n(20) Chapman, D. J. Am. Oil Chem. Soc. 1965, 42, 353-371."])
    assert bibliography_element("(20) Chapman, D. J. Am. Oil Chem. Soc. 1965, 42, 353-371.", records)


def test_image_reference_numbers_do_not_guess_an_original_bibliography():
    from afs_assistant.core.models import ExtractedDocument
    d = ExtractedDocument(source=source(), parser_version="test", processing_key="key", markdown="",
        page_count=1, elements=[])
    OfficialExtractor._image_references(d, "glucose [110] and vitamins [132]", 17)
    assert {r.label for r in d.references} == {"figure-p17-110", "figure-p17-132"}
    assert all(r.status == "unresolved_marker" and r.target_source_id is None for r in d.references)


def test_workbook_rows_preserve_units_notes_and_real_locations(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "table.xlsx"
    w = Workbook(); ws = w.active
    assert ws is not None
    ws.title = "method_comparison"
    ws.append(["method_no","method","measurement_technique","physical_phenomenon","sample_type",
        "optical_path_or_depth","sensitivity","selectivity","cost"])
    for n in range(19):
        ws.append([n+1, "NIR", "Transmission", "absorption", "liquid", "1-5 mm", "average", "average", "average"])
    w.create_sheet("notes").append(["Qualitative ratings, not measured accuracy."])
    w.save(path); w.close()
    s = source("AFS_TABLE", format="xlsx")
    e = OfficialExtractor(Settings(data_dir=tmp_path))
    d = e.extract(SourceMaterial(source=s, path=path), tmp_path / "parsed")
    batch = OfficialChunker(Settings(chunk_size=128)).process_material(d)
    assert len(batch.items) == 19
    item = batch.items[0]
    assert item.cell_range == "A2:I2" and item.provenance == []
    assert item.citation_label == "S01 Table 1 (p.6)"
    assert item.related_source_id == "S01" and item.related_page == 6
    assert item.structured_data["optical_path_or_depth"] == "1-5 mm"
    # Notes are attached once as qualifications rather than embedded in every row.
    assert "Qualitative ratings" not in item.text
    assert item.qualifications == ["Qualitative ratings, not measured accuracy."]


def slide_document(page, ocr):
    from afs_assistant.core.models import DocumentElement, ExtractedDocument, Provenance
    s = source("S02", visual_strategy="page")
    where = [Provenance(page=page)]
    return ExtractedDocument(source=s, parser_version="test", processing_key="key", markdown="", page_count=20, elements=[
        DocumentElement(element_id=f"page-{page}-heading", label="section_header", heading_level=1, text=f"Slide {page}: Fibers", provenance=where),
        DocumentElement(element_id=f"page-{page}-ocr", label="text", text=ocr, provenance=where),
        DocumentElement(element_id=f"page-{page}-visual", label="picture", text=f"Complete slide {page}", caption=f"Complete slide {page}",
                        notes=ocr, image=f"element-{900000 + page}.png", provenance=where)])


def test_curated_transcription_replaces_scrambled_slide_ocr():
    ocr = "Polycrystalline\n3000 nm\n17μm\nInfrared Fibers\nStandard VIS-NIR Silica Fibers\n240rm"
    batch = OfficialChunker(Settings(chunk_size=128)).process_material(slide_document(6, ocr))
    texts = [item for item in batch.items if item.element_type == "text"]
    assert not any("240rm" in item.text for item in texts)
    curated = [item for item in texts if item.element_ids[0].startswith("curated-") and item.locator.label == "PDF p. 6"]
    assert len(curated) == 1 and curated[0].provenance[0].page == 6
    assert "Polycrystalline infrared fibers: 3000 nm to 17 µm" in curated[0].text
    assert "350 nm to 2400 nm" in curated[0].text
    image = next(item for item in batch.items if item.image)
    # Vision context keeps the OCR; the indexed slide image text does not repeat it.
    assert image.image is not None and "240rm" in image.image.nearby_text
    assert "240rm" not in image.text


def test_slides_without_transcription_keep_ocr_and_drop_heading_only_chunks():
    batch = OfficialChunker(Settings(chunk_size=128)).process_material(slide_document(5, "Raman FTIR Fluorescence NIR modular system"))
    texts = [item.text for item in batch.items if item.element_type == "text"]
    assert any("modular system" in text for text in texts)
    assert not any(item.element_ids == ["page-5-heading"] for item in batch.items)


def test_branding_and_credit_lines_are_not_indexed():
    from afs_assistant.core.models import DocumentElement, ExtractedDocument, Provenance
    where = [Provenance(page=10)]
    d = ExtractedDocument(source=source("S01"), parser_version="test", processing_key="key", markdown="", page_count=19, elements=[
        DocumentElement(element_id="h", label="section_header", heading_level=1, text="5.2 Near-IR", provenance=where),
        DocumentElement(element_id="t1", label="text", text="NIR spectra (780-2500 nm) contain overtones of fundamental vibrations.", provenance=where),
        DocumentElement(element_id="h2", label="section_header", heading_level=1, text="Fig. 2", provenance=where),
        DocumentElement(element_id="t2", label="text", text="Image courtesy of art photonics GmbH.\nMiR PHOTONICS", provenance=where)])
    texts = [item.text for item in OfficialChunker(Settings(chunk_size=128)).process_material(d).items]
    assert any("780-2500 nm" in text for text in texts)
    assert not any("courtesy" in text or "PHOTONICS" in text for text in texts)


def test_short_webpage_fragments_are_joined():
    from afs_assistant.core.models import DocumentElement, ExtractedDocument
    d = ExtractedDocument(source=source("S06", format="html"), parser_version="test", processing_key="key", markdown="",
        page_count=0, elements=[DocumentElement(element_id=f"html-{n}", label="text", text=t)
        for n, t in enumerate(["Laboratory", "Spectroscopy", "We test customer processes with fiber probes."])])
    items = OfficialChunker(Settings(chunk_size=128)).process_material(d).items
    assert len(items) == 1 and "Laboratory" in items[0].text and "fiber probes" in items[0].text
