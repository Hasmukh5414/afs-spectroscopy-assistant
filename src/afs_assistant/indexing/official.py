"""Admission and format-aware preparation of the supplied AFS knowledge bundle."""
from dataclasses import replace
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import shutil

from afs_assistant.core.fingerprints import fingerprint
from afs_assistant.core.models import (
    ChunkBatch, DocumentElement, EvidenceItem, ExtractedDocument, Locator,
    Provenance, ReferenceRecord,
)
from .chunking import DocumentChunker
from .extractor import DoclingExtractor
from .material_cache import write_json
from .references import bibliography_element, extract_references, normalized


OFFICIAL_FILES = [
    ("S01", "S01_Multispectral_Laboratory.pdf", "Multispectral Laboratory", "technical", "regions", []),
    ("S02", "S02_AFS_Process_Control_Slides.pdf", "AFS Process Control Slides", "technical", "page", []),
    ("S03", "S03_LED_Kidney_Study_2017.pdf", "Development and Testing of an LED-Based Near-Infrared Sensor", "research", "regions",
     ["10.3390/s17081914", "Sensors 2017, 17, 1914"]),
    ("S04", "S04_Combi_MIR_Fluorescence_2021.pdf", "Fiber Probe for Simultaneous Mid-Infrared and Fluorescence Spectroscopic Analysis", "research", "regions",
     ["10.1021/acs.analchem.1c00080"]),
    ("S05", "S05_AFS_Flyer.pdf", "AFS Flyer", "business", "page", []),
    ("AFS_TABLE", "AFS_method_comparison_table.xlsx", "AFS Method Comparison Table", "technical", "regions", []),
]


def admit_official(materials_dir: Path, data_dir: Path, *, snapshot_web=True) -> Path:
    """Copy only knowledge sources; scenario text and instructions never enter the index."""
    from pypdf import PdfReader
    root = materials_dir.resolve()
    data_dir = data_dir.resolve()
    records = []
    for sid, filename, title, role, visual, keys in OFFICIAL_FILES:
        original = root / filename
        content = original.read_bytes()
        digest = hashlib.sha256(content).hexdigest()
        retained = data_dir / "originals" / digest / filename
        retained.parent.mkdir(parents=True, exist_ok=True)
        if retained.exists() and hashlib.sha256(retained.read_bytes()).hexdigest() != digest:
            raise ValueError("Retained source identity mismatch")
        if not retained.exists():
            shutil.copyfile(original, retained)
        fmt = "pdf" if filename.endswith(".pdf") else "xlsx"
        records.append({"source_id": sid, "title": title, "status": "acquired_unreviewed",
            "path": str(retained.relative_to(data_dir)), "sha256": digest, "url": f"urn:afs:{sid}",
            "role": "guidance" if role == "business" else "application",
            "format": fmt, "authority": "provided", "evidence_role": role,
            "visual_strategy": visual, "citation_keys": keys,
            "structured_table_pages": [6] if sid == "S01" else [],
            "expected_pages": len(PdfReader(original).pages) if fmt == "pdf" else 1})
    if snapshot_web:
        import httpx
        link_text = (root / "S06_Laboratory_Link.txt").read_text()
        match = re.search(r"https://[^\s<>]+", link_text)
        if not match or match[0].rstrip("/").lower() != "https://mirphotonics.com/laboratory":
            raise ValueError("The supplied laboratory link is not the expected public source")
        url = match[0]
        previous = data_dir / "originals/official-materials.json"
        cached = next((row for row in json.loads(previous.read_text()) if row.get("source_id") == "S06"), None) if previous.is_file() else None
        cached_path = (data_dir / cached["path"]).resolve() if cached else None
        if cached and cached_path and cached_path.is_relative_to(data_dir) and cached_path.is_file() and hashlib.sha256(cached_path.read_bytes()).hexdigest() == cached["sha256"] and cached["url"] == url:
            content = cached_path.read_bytes()
        else:
            cached = None
            with httpx.Client(follow_redirects=True, timeout=45) as client:
                response = client.get(url)
                response.raise_for_status()
            content = response.content
        if not content or len(content) > 10_000_000:
            raise ValueError("Laboratory snapshot is empty or exceeds the preparation limit")
        digest = hashlib.sha256(content).hexdigest()
        retained = data_dir / "originals" / digest / "S06-laboratory.html"
        retained.parent.mkdir(parents=True, exist_ok=True)
        retained.write_bytes(content)
        records.append({"source_id": "S06", "title": "AFS Laboratory (supplied webpage)",
            "status": "acquired_unreviewed", "path": str(retained.relative_to(data_dir)),
            "sha256": digest, "url": url, "role": "guidance", "format": "html",
            "authority": "provided", "evidence_role": "business", "expected_pages": 1,
            "snapshot_at": cached["snapshot_at"] if cached else datetime.now(timezone.utc).isoformat()})
    # Organizer guidance is kept separately from the scientific evidence corpus.
    from zipfile import ZipFile
    from xml.etree import ElementTree as ET
    guide_path = root / "AFS_Examples_and_Reference.docx"
    if guide_path.is_file():
        with ZipFile(guide_path) as archive:
            xml = ET.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        paragraphs = ["".join(t.text or "" for t in p.findall(".//w:t", ns)) for p in xml.findall(".//w:p", ns)]
        # The heading also appears in the table of contents. Start at its last
        # occurrence so scenarios and worked answers cannot enter runtime guidance.
        start = max(i for i, text in enumerate(paragraphs) if text.startswith("3. Answer structure"))
        write_json(data_dir / "guidance/organizer-guide.json", {
            "source": guide_path.name, "sha256": hashlib.sha256(guide_path.read_bytes()).hexdigest(),
            "role": "response_and_source_navigation_guidance_not_scientific_evidence",
            "text": "\n".join(paragraphs[start:])})
    receipt = data_dir / "originals" / "official-materials.json"
    write_json(receipt, records)
    return receipt


class OfficialExtractor:
    """Use Docling's integrated OCR, retained visuals and bibliography metadata."""
    VERSION = "afs-official-v2-docling-ocr"

    def __init__(self, settings):
        self.settings = settings
        policy = replace(settings, do_ocr=True, image_scale=2.0)
        self.pdf = DoclingExtractor(policy, ocr_mode="pdf_aware_layout_regions")
        self.slides = DoclingExtractor(policy, ocr_mode="full_page")
        self.parser_version = self.pdf.parser_version

    def processing_key(self, source_hash):
        return fingerprint({"source": source_hash, "official": self.VERSION,
            "docling": self.pdf.processing_key(source_hash),
            "slides": self.slides.processing_key(source_hash), "ocr": "docling-integrated",
            "visual_policy": "whole-slide-and-technical-regions-v1"})

    def extract(self, material, directory):
        directory.mkdir(parents=True, exist_ok=True)
        source = material.source
        if source.format == "xlsx":
            return self._workbook(material)
        if source.format == "html":
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(material.path.read_bytes(), "html.parser")
            for node in soup(["script", "style", "nav", "footer", "header"]):
                node.decompose()
            container = soup.find("main") or soup.find("article") or soup
            elements = [DocumentElement(element_id=f"html-{n}", label="text", text=p.get_text(" ", strip=True))
                for n, p in enumerate(container.find_all(["h1", "h2", "h3", "p", "li"]))
                if p.get_text(" ", strip=True)]
            if not elements:
                raise ValueError("The supplied webpage contains no readable material")
            return ExtractedDocument(source=source, parser_version="beautifulsoup",
                processing_key=self.processing_key(source.sha256), markdown="\n\n".join(e.text for e in elements),
                page_count=0, elements=elements)
        from pypdf import PdfReader
        from PIL import Image
        parser = self.slides if source.visual_strategy == "page" else self.pdf
        document = parser.extract(material, directory)
        pages = [p.extract_text() or "" for p in PdfReader(material.path).pages]
        document.processing_key = self.processing_key(source.sha256)
        document.references = extract_references(source, pages)
        kept = []
        seen_text = set()
        seen_images = set()
        located_text = {}
        for element in document.elements:
            if element.text.strip() and not bibliography_element(element.text, document.references):
                for location in element.provenance:
                    located_text.setdefault(location.page, []).append(element.text)
        for element in document.elements:
            if element.label in {"page_header", "page_footer"}:
                continue
            if bibliography_element(element.text, document.references):
                continue
            if source.visual_strategy == "page":
                continue
            if element.label == "picture" and element.image:
                path = directory / element.image
                pixels = path.read_bytes()
                identity = hashlib.sha256(pixels).hexdigest()
                # Repeated page logos do not need independent image descriptions.
                if identity in seen_images:
                    continue
                seen_images.add(identity)
                with Image.open(path) as image:
                    width, height = image.size
                if width < 90 or height < 55:
                    continue
            elif element.text.strip():
                identity = (element.provenance[0].page, normalized(element.text))
                if identity in seen_text:
                    continue
                seen_text.add(identity)
            kept.append(element)
        for page_no, native in enumerate(pages, 1):
            page_image = directory / document.page_images[page_no]
            ocr = document.ocr_text.get(page_no, "")
            self._image_references(document, ocr, page_no)
            if source.visual_strategy == "page":
                # Docling owns both OCR and PDF text recovery; no second OCR pass.
                body = "\n".join(located_text.get(page_no, []))
                present = normalized(body)
                recovered = [line for line in ocr.splitlines()
                    if len(line.strip()) > 2 and normalized(line) not in present]
                body = "\n".join([body, *recovered]).strip() or native
                title = next((line for line in body.splitlines() if len(line.strip()) > 12), source.title)[:180]
                provenance = [Provenance(page=page_no)]
                kept.append(DocumentElement(element_id=f"page-{page_no}-heading", label="section_header",
                    heading_level=1, text=f"Slide {page_no}: {title}", provenance=provenance))
                kept.append(DocumentElement(element_id=f"page-{page_no}-ocr", label="text", text=body, provenance=provenance))
                image_name = f"element-{900000 + page_no}.png"
                shutil.copyfile(page_image, directory / image_name)
                kept.append(DocumentElement(element_id=f"page-{page_no}-visual", label="picture",
                    text=f"Complete slide {page_no}: {title}", caption=f"Complete slide {page_no}: {title}",
                    notes=body, image=image_name, provenance=provenance))
                document.coverage.append({"source_id": source.source_id, "page": page_no,
                    "native_chars": len(native), "ocr_chars": len(ocr), "extracted_chars": len(body),
                    "ocr_engine": "docling", "ocr_mode": "full_page",
                    "visual_strategy": "page", "image_retained": True})
                if len(body.strip()) < 25:
                    raise ValueError(f"Page {page_no} has no usable native or OCR content")
            else:
                present = normalized("\n".join(located_text.get(page_no, [])))
                recovered = [line for line in ocr.splitlines()
                    if len(line.strip()) > 2 and normalized(line) not in present]
                if recovered:
                    kept.append(DocumentElement(element_id=f"page-{page_no}-ocr-labels", label="text",
                        text="\n".join(recovered), provenance=[Provenance(page=page_no)]))
                document.coverage.append({"source_id": source.source_id, "page": page_no,
                    "native_chars": len(native), "ocr_chars": len(ocr),
                    "ocr_engine": "docling", "ocr_mode": "pdf_aware_layout_regions",
                    "visual_strategy": "regions", "image_retained": True})
        document.elements = kept
        document.markdown = "\n\n".join(e.text for e in kept if e.text)
        write_json(directory / "reference-registry.json", [r.model_dump() for r in document.references])
        write_json(directory / "coverage.json", document.coverage)
        return document

    @staticmethod
    def _image_references(document, text, page):
        known = {r.label for r in document.references}
        for number in set(re.findall(r"\[(\d{1,3})\]", text)):
            if number not in known:
                label = f"figure-p{page}-{number}"
                if label not in known:
                    document.references.append(ReferenceRecord(source_id=document.source.source_id,
                        label=label, citation=f"Unresolved printed reference marker [{number}] inside an image",
                        location=f"p. {page}", status="unresolved_marker"))
                    known.add(label)

    def _workbook(self, material):
        from openpyxl import load_workbook
        workbook = load_workbook(material.path, read_only=True, data_only=True)
        try:
            sheet = workbook["method_comparison"]
            rows = list(sheet.iter_rows(values_only=True))
            headers = [str(v) for v in rows[0]]
            notes = [str(row[0]) for row in workbook["notes"].iter_rows(values_only=True) if row and row[0]]
            elements = []
            for number, values in enumerate(rows[1:], 2):
                if not any(v is not None for v in values):
                    continue
                data = dict(zip(headers, values))
                text = "\n".join(f"{key}: {value}" for key, value in data.items() if value is not None)
                elements.append(DocumentElement(element_id=f"row-{number}", label="table",
                    text=text, notes="\n".join(notes), table_data={"row": number, "values": data,
                    "notes": notes, "worksheet": sheet.title, "cell_range": f"A{number}:I{number}"}))
            if len(elements) != 19:
                raise ValueError("Official comparison workbook row coverage differs from the supplied 19 rows")
            return ExtractedDocument(source=material.source, parser_version="openpyxl",
                processing_key=self.processing_key(material.source.sha256),
                markdown="\n\n".join(e.text for e in elements), page_count=0, elements=elements)
        finally:
            workbook.close()

    @staticmethod
    def validate_corpus(corpus):
        sources = {s.source_id: s for s in corpus.sources}
        required = {row[0] for row in OFFICIAL_FILES} | {"S06"}
        if not required.issubset(sources):
            raise ValueError("An official release is missing required supplied sources")
        coverage = corpus.processing.get("coverage", [])
        if len(coverage) != 64:
            raise ValueError("Official PDF coverage must account for all 64 pages")
        slides = [row for row in coverage if row["source_id"] == "S02"]
        if len(slides) != 20 or any(row.get("extracted_chars", 0) < 25 for row in slides):
            raise ValueError("A supplied slide has insufficient recovered text")
        if not any(row["page"] == 20 and row.get("ocr_chars", 0) > 300 for row in slides):
            raise ValueError("The image-only infographic was not recovered by Docling OCR")
        table = [i for i in corpus.items if i.source_id == "AFS_TABLE"]
        if len(table) != 19 or any(not i.cell_range for i in table):
            raise ValueError("Workbook row provenance is incomplete")
        for sid, count in (("S03", 43), ("S04", 25)):
            refs = [r for r in corpus.references if r.source_id == sid and r.label.isdigit()]
            if len(refs) != count:
                raise ValueError(f"Reference coverage is incomplete for {sid}")
        if not any(r.source_id == "S04" and r.label == "14" and r.target_source_id == "S03" for r in corpus.references):
            raise ValueError("The supplied-paper reference link S04:14 to S03 is missing")


CURATED_PATH = Path(__file__).with_name("curated_transcriptions.json")
CLEANUP_VERSION = "afs-official-cleanup-v1"
# Repeated logos, credits and running headers carry no technical content.
BOILERPLATE_LINES = re.compile(r"(?:images? courtesy of art photonics gmbh\.?|mi?r\W?|mir photonics|photonics|"
    r"art photonics|member of the nynomic group|www\.artphotonics\.com|www\.afs\.art|pubs\.acs\.org/ac technical note)",
    re.IGNORECASE)


def curated_transcriptions() -> dict:
    return json.loads(CURATED_PATH.read_text(encoding="utf-8"))


class OfficialChunker(DocumentChunker):
    def __init__(self, settings):
        super().__init__(settings)
        self.curated = curated_transcriptions()
        self.cache_identity = fingerprint({"chunker": self.cache_identity, "cleanup": CLEANUP_VERSION,
            "curated": self.curated})

    def process_material(self, document):
        if document.source.format == "pdf":
            # Also apply bibliography filtering when reusing an extraction cache.
            filtered = document.model_copy(update={"elements": [
                element for element in document.elements
                if not bibliography_element(element.text, document.references)
                and not (element.provenance and all(p.page in document.source.structured_table_pages for p in element.provenance))]})
            return self._clean(super().process_material(filtered), document)
        items = []
        source = document.source
        elements = document.elements
        if source.format == "html":
            # Webpage paragraphs are short fragments; join them into sentence-bounded chunks.
            elements = [element.model_copy(update={"text": "\n".join(e.text for e in elements)}) for element in elements[:1]]
        for element in elements:
            if source.format == "xlsx":
                # Workbook notes stay attached as qualifications instead of repeating in every embedded row.
                pieces = [element.text]
            else:
                pieces, current = [], []
                for sentence in self._sentences(element.text):
                    if current and self.count("\n".join(current + [sentence])) > self.settings.chunk_size:
                        pieces.append("\n".join(current))
                        current = []
                    current.append(sentence)
                if current:
                    pieces.append("\n".join(current))
            for number, text in enumerate(pieces):
                identity = fingerprint([source.sha256, element.element_id, number, text, self.settings.chunk_size])
                row = element.table_data
                items.append(EvidenceItem(evidence_id=f"{source.source_id}_{identity[:40]}",
                    source_id=source.source_id, source_version=source.sha256, title=source.title, text=text,
                    locator=Locator(kind="row" if row else "section",
                        label=f"{row['worksheet']}!{row['cell_range']}" if row else "Supplied laboratory webpage snapshot"),
                    kind="method_comparison" if row else "guidance", synthetic=False,
                    element_type="table_row" if row else "text", element_ids=[element.element_id],
                    worksheet=row.get("worksheet", ""), cell_range=row.get("cell_range", ""),
                    structured_data=row.get("values", {}), qualifications=row.get("notes", []),
                    citation_label="S01 Table 1 (p.6)" if row else "",
                    related_source_id="S01" if row else "", related_page=6 if row else None,
                    oversized=self.count(text) > self.settings.chunk_size))
        return ChunkBatch(items=items, contexts=[])

    def _clean(self, batch, document):
        """Replace unreadable extractions with curated transcriptions and drop non-content chunks."""
        source = document.source
        entries = [e for e in self.curated["entries"] if e["source_id"] == source.source_id]
        replaced_ids = {eid for e in entries for eid in e.get("replaces_element_ids", [])}
        replaced_table_pages = {e["page"] for e in entries if e.get("replaces_tables_on_page")}
        headings = {e.element_id for e in document.elements if e.label in {"title", "section_header"}}
        items = []
        for item in batch.items:
            pages = {p.page for p in item.provenance}
            if item.element_type == "table_row" and pages & replaced_table_pages:
                continue
            if item.element_type == "text":
                if replaced_ids & set(item.element_ids) or set(item.element_ids) <= headings:
                    continue
                kept = [line for line in item.text.splitlines() if not BOILERPLATE_LINES.fullmatch(line.strip())]
                text = "\n".join(kept).strip()
                if len(re.findall(r"[A-Za-z]{2,}", text)) < 3:
                    continue
                if text != item.text:
                    item = item.model_copy(update={"text": text, "evidence_id": source.source_id[:40] + "_" +
                        fingerprint([item.evidence_id, text])[:40], "oversized": self.count(text) > self.settings.chunk_size})
            items.append(item)
        for entry in entries:
            items.append(EvidenceItem(
                evidence_id=f"{source.source_id[:40]}_curated_" + fingerprint([self.curated["version"], entry])[:40],
                source_id=source.source_id, source_version=source.sha256, title=source.title, text=entry["text"],
                locator=Locator(kind="page", label=f"PDF p. {entry['page']}"), kind=source.kind, synthetic=False,
                qualifications=[self.curated["policy"]], element_ids=[f"curated-{entry['id']}"],
                provenance=[Provenance(page=entry["page"])], heading_path=[entry["title"]],
                element_type=entry["element_type"], issue_ids=source.issue_ids,
                oversized=self.count(entry["text"]) > self.settings.chunk_size))
        return ChunkBatch(items=items, contexts=batch.contexts)


def prepare_official(settings, store, gateway, materials_dir, release, *, embed=False, describe=False,
                     publish_release=False):
    from .builder import EvidenceBuilder
    if publish_release and not (embed and describe):
        raise ValueError("Official publication requires both image descriptions and embeddings.")
    receipt = admit_official(Path(materials_dir), settings.data_dir)
    policy = replace(settings, chunk_size=128)
    builder = EvidenceBuilder(policy, store, gateway,
        extractor=OfficialExtractor(policy), chunker=OfficialChunker(policy))
    report = builder.process_materials(receipt, release, generate_embeddings=embed,
        describe_images=describe, publish_release=publish_release)
    return report
